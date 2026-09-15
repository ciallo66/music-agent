"""更新本地简历文档中的项目经历内容。"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

RESUME = Path(r"C:\Users\28265\Desktop\面试.docx")
OUTPUT = Path(r"C:\Users\28265\Desktop\面试_更新版.docx")


def clear_runs(paragraph):
    """清除段落中的现有文本片段。"""
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def add_run(paragraph, text, *, size=10.0, bold=False, color="000000"):
    """向段落添加指定字体样式的文本片段。"""
    run = paragraph.add_run(text)
    run.font.name = "等线"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def replace_bullet(paragraph, label, body):
    """用带圆点的标签和正文替换段落内容。"""
    clear_runs(paragraph)
    add_run(paragraph, "●", size=9.0, color="0070C0")
    add_run(paragraph, "  ", size=9.0, color="0070C0")
    add_run(paragraph, label, size=10.0, bold=True, color="1F3448")
    add_run(paragraph, body, size=10.0, color="000000")


def main():
    """更新目标简历并保存到输出路径。"""
    doc = Document(str(RESUME))
    paragraphs = doc.paragraphs

    internship = [
        ("功能定位：", "将银行流水 PDF/图片整理为可复核的结构化 Excel。"),
        ("解析链路：", "负责原生文本与 PaddleOCR 识别，完成页面渲染和交易记录提取。"),
        ("版式适配：", "处理多银行表头差异、跨页记录和金额/余额错位。"),
        ("质量闭环：", "保留页码、置信度和异常原因，校验余额/主体/重复记录并输出复核结果。"),
    ]
    music = [
        ("产品功能：", "面向多用户提供歌曲搜索、播放、歌词、歌单和个性化推荐。"),
        ("前端实现：", "负责 Vue3 页面与播放器状态，串起搜索、详情、收藏和歌单操作。"),
        (
            "AI 助手：",
            "用 FastAPI + DeepSeek Function Calling 支持搜歌、歌曲分析、相似推荐和口味问答。",
        ),
        ("推荐检索：", "基于播放/收藏行为生成推荐，使用 pgvector 完成相似歌曲与音乐知识检索。"),
        ("数据部署：", "接入 Jamendo 音乐数据，完成 SSE 流式响应和 Docker + Nginx 部署。"),
    ]
    interview = [
        ("产品功能：", "面向刷题用户提供题目浏览、筛选、收藏、登录和搜索。"),
        ("全栈实现：", "负责 Vue3 页面、Express5 接口和 MySQL 数据设计，打通题库与用户数据。"),
        ("安全上线：", "使用 JWT、bcrypt 和 DOMPurify 完成鉴权与输入防护，并部署到 Vercel/ECS。"),
    ]

    for index, (label, body) in zip(range(14, 18), internship, strict=False):
        replace_bullet(paragraphs[index], label, body)
    for index, (label, body) in zip(range(21, 25), music[:4], strict=False):
        replace_bullet(paragraphs[index], label, body)
    music_tail = paragraphs[25].insert_paragraph_before()
    replace_bullet(music_tail, *music[4])
    for index, (label, body) in zip(range(27, 30), interview, strict=False):
        replace_bullet(paragraphs[index], label, body)

    doc.save(str(OUTPUT))


if __name__ == "__main__":
    main()
