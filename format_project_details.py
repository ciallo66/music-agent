"""格式化面试文档中的项目经历段落。"""

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

path = r"C:\Users\28265\Desktop\面试.docx"
doc = Document(path)
BLACK = "000000"
NAVY = "1F3448"
BLUE = "0070C0"


def clear(p):
    """清除段落中的现有文字节点。"""
    for child in list(p._element):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            p._element.remove(child)


def fmt(r, size=10.0, bold=False, color=BLACK):
    """设置文本字体、字号、粗细和颜色。"""
    r.font.name = "微软雅黑"
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)


def bullet(p, label, body):
    """将段落格式化为带标签和正文的项目要点。"""
    clear(p)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("●  ")
    fmt(r, 9, False, BLUE)
    r = p.add_run(label + "：")
    fmt(r, 10.0, True, NAVY)
    r = p.add_run(body)
    fmt(r, 10.0, False, BLACK)


paras = doc.paragraphs
ai = [
    ("前端", "Vue3、TypeScript、Pinia、Vite 搭建音乐库、播放器、歌单、歌词和个人画像页面。"),
    (
        "Agent / AI",
        "DeepSeek API、Function Calling、SSE 实现歌曲搜索、音乐分析、相似歌曲和偏好分析助手。",
    ),
    ("RAG / 向量检索", "pgvector 与 Embedding 支持相似歌曲检索和音乐知识问答。"),
    (
        "后端 / 数据库 / 部署",
        "FastAPI、SQLAlchemy、PostgreSQL 完成用户、歌曲和推荐接口，接入 Jamendo 数据导入，"
        "并用 Docker Compose、Nginx 部署。",
    ),
]
for paragraph, (label, body) in zip(paras[21:25], ai, strict=False):
    bullet(paragraph, label, body)
old = [
    ("前端", "Vue3、Vue Router、Pinia、Element Plus 实现题目浏览、筛选、收藏和登录状态管理。"),
    ("后端", "Node.js、Express5 提供登录、题目、标签等 RESTful 接口。"),
    ("数据库", "MySQL 全文索引和 ngram 分词实现中英文搜索，多对多标签和事务保证数据一致。"),
    (
        "安全 / 部署",
        "JWT、bcrypt、Zod、参数化查询和 DOMPurify 完成鉴权与输入安全，"
        "GitHub Actions、Vercel、ECS、Nginx 完成部署。",
    ),
]
for paragraph, (label, body) in zip(paras[27:31], old, strict=False):
    bullet(paragraph, label, body)
doc.save(path)
